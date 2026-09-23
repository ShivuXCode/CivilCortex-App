import io
import re
from fpdf import FPDF
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from app.models.inspection import Inspection, InspectionImage
from app.models.defect import Assessment
from app.models.hierarchy import Building
from app.services.storage_service import StorageService

class DocumentService:
    def __init__(self):
        self.storage_service = StorageService()

    def _get_image_bytes(self, image: InspectionImage) -> bytes:
        if not image:
            return b""
        try:
            return self.storage_service.get_file_bytes(image.file_path)
        except Exception:
            return b""

    def generate_pdf_report(self, inspection: Inspection, building: Building, assessment: Assessment, image: InspectionImage) -> bytes:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_auto_page_break(auto=True, margin=15)
        
        # Title
        pdf.set_font("helvetica", "B", 18)
        pdf.cell(0, 10, "Engineering Assessment Report", ln=True, align="C")
        pdf.ln(5)
        
        # Metadata
        pdf.set_font("helvetica", "", 10)
        pdf.cell(0, 6, f"Inspection ID: {inspection.id}", ln=True)
        pdf.cell(0, 6, f"Generated At: {assessment.updated_at.strftime('%Y-%m-%d %H:%M:%S UTC')}", ln=True)
        pdf.ln(5)
        
        # Building Details
        pdf.set_font("helvetica", "B", 14)
        pdf.cell(0, 8, "Building Details", ln=True)
        pdf.set_font("helvetica", "", 11)
        pdf.cell(0, 6, f"Name: {building.name}", ln=True)
        pdf.cell(0, 6, f"Location: {getattr(building, 'location', 'N/A')}", ln=True)
        pdf.ln(5)
        
        # Image
        if image:
            img_bytes = self._get_image_bytes(image)
            if img_bytes:
                # FPDF2 allows passing io.BytesIO to image()
                try:
                    pdf.set_font("helvetica", "B", 14)
                    pdf.cell(0, 8, "Original Inspection Image", ln=True)
                    img_stream = io.BytesIO(img_bytes)
                    pdf.image(img_stream, w=150)
                    pdf.ln(5)
                except Exception:
                    pass

        # Defect & Assessment
        pdf.set_font("helvetica", "B", 14)
        pdf.cell(0, 8, "Defect & Assessment", ln=True)
        pdf.set_font("helvetica", "", 11)
        pdf.cell(0, 6, f"Severity: {assessment.severity}", ln=True)
        pdf.cell(0, 6, f"Risk Level: {assessment.risk}", ln=True)
        pdf.cell(0, 6, f"Priority: {assessment.priority}", ln=True)
        pdf.ln(5)
        
        # Engineer Review
        if inspection.status in ["APPROVED", "COMPLETED"]:
            pdf.set_font("helvetica", "B", 14)
            pdf.cell(0, 8, "Engineer Action Plan", ln=True)
            pdf.set_font("helvetica", "", 11)
            pdf.multi_cell(0, 6, assessment.repair_recommendation or "No action plan recorded.")
            pdf.ln(5)
        
        # RAG / Regulatory Evidence
        pdf.set_font("helvetica", "B", 14)
        pdf.cell(0, 8, "Regulatory / Standard Evidence", ln=True)
        pdf.set_font("helvetica", "", 11)
        pdf.multi_cell(0, 6, assessment.rag_context or "No specific standards retrieved.")
        pdf.ln(5)
        
        # Executive Report
        pdf.set_font("helvetica", "B", 14)
        pdf.cell(0, 8, "Executive Report", ln=True)
        pdf.set_font("helvetica", "", 11)
        if assessment.llm_report:
            # FPDF2 multi_cell supports basic markdown
            try:
                pdf.multi_cell(0, 6, assessment.llm_report, markdown=True)
            except Exception:
                # Fallback if markdown parsing fails
                pdf.multi_cell(0, 6, assessment.llm_report)
        else:
            pdf.multi_cell(0, 6, "Report not yet generated.")
            
        return bytes(pdf.output())
        
    def _add_markdown_to_docx(self, doc, text: str):
        lines = text.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('- '):
                p = doc.add_paragraph(style='List Bullet')
                self._add_inline_markdown(p, line[2:])
            else:
                p = doc.add_paragraph()
                self._add_inline_markdown(p, line)
                
    def _add_inline_markdown(self, paragraph, text: str):
        # Very simple bold parsing for **text**
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            else:
                paragraph.add_run(part)

    def generate_docx_report(self, inspection: Inspection, building: Building, assessment: Assessment, image: InspectionImage) -> bytes:
        doc = Document()
        
        # Title
        title = doc.add_heading('Engineering Assessment Report', 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Metadata
        doc.add_paragraph(f"Inspection ID: {inspection.id}")
        doc.add_paragraph(f"Generated At: {assessment.updated_at.strftime('%Y-%m-%d %H:%M:%S UTC')}")
        
        # Building Details
        doc.add_heading('Building Details', level=1)
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = 'Name'
        table.cell(0, 1).text = building.name or "N/A"
        table.cell(1, 0).text = 'Location'
        table.cell(1, 1).text = getattr(building, 'location', 'N/A') or "N/A"
        
        # Image
        if image:
            img_bytes = self._get_image_bytes(image)
            if img_bytes:
                doc.add_heading('Original Inspection Image', level=1)
                try:
                    img_stream = io.BytesIO(img_bytes)
                    doc.add_picture(img_stream, width=Inches(5))
                except Exception:
                    doc.add_paragraph("(Failed to embed image)")
                    
        # Defect & Assessment
        doc.add_heading('Defect & Assessment', level=1)
        doc.add_paragraph(f"Severity: {assessment.severity}")
        doc.add_paragraph(f"Risk Level: {assessment.risk}")
        doc.add_paragraph(f"Priority: {assessment.priority}")
        
        # Engineer Review
        if inspection.status in ["APPROVED", "COMPLETED"]:
            doc.add_heading('Engineer Action Plan', level=1)
            doc.add_paragraph(assessment.repair_recommendation or "No action plan recorded.")
            
        # Regulatory Evidence
        doc.add_heading('Regulatory / Standard Evidence', level=1)
        doc.add_paragraph(assessment.rag_context or "No specific standards retrieved.")
        
        # Executive Report
        doc.add_heading('Executive Report', level=1)
        if assessment.llm_report:
            self._add_markdown_to_docx(doc, assessment.llm_report)
        else:
            doc.add_paragraph("Report not yet generated.")
            
        # Save to bytes
        file_stream = io.BytesIO()
        doc.save(file_stream)
        return file_stream.getvalue()
